"""Multi-format parsers producing a single normalized representation."""

from __future__ import annotations

import html as _html
import re
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class NormalizedDocument:
    doc_id: str
    title: str
    text: str
    source_path: str
    extension: str
    category: str
    raw_metadata: dict = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


def _read_text_file(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def parse_markdown(path: Path, text: str | None = None) -> tuple[str, dict]:
    raw = text if text is not None else _read_text_file(path)
    meta: dict = {}
    # Optional YAML front-matter
    if raw.startswith("---"):
        end = raw.find("\n---", 3)
        if end != -1:
            front = raw[3:end].strip()
            raw = raw[end + 4:].lstrip()
            for line in front.splitlines():
                if ":" in line:
                    key, _, value = line.partition(":")
                    meta[key.strip().lower()] = value.strip().strip("'\"")
    # First ATX heading becomes title candidate
    title = meta.get("title", "")
    if not title:
        match = re.search(r"^#{1,6}\s+(.+)$", raw, re.MULTILINE)
        if match:
            title = match.group(1).strip()
    if title:
        meta["title"] = title
    return raw, meta


def parse_txt(path: Path) -> tuple[str, dict]:
    raw = _read_text_file(path)
    lines = [line.strip() for line in raw.splitlines() if line.strip()]
    meta: dict = {}
    if lines:
        meta["title"] = lines[0][:120]
    return raw, meta


def parse_html(path: Path, text: str | None = None) -> tuple[str, dict]:
    raw = text if text is not None else _read_text_file(path)
    meta: dict = {}
    title_match = re.search(r"<title[^>]*>(.*?)</title>", raw, re.IGNORECASE | re.DOTALL)
    if title_match:
        meta["title"] = _html.unescape(re.sub(r"\s+", " ", title_match.group(1)).strip())
    # Strip scripts/styles, then tags
    cleaned = re.sub(r"(?is)<(script|style).*?</\1>", " ", raw)
    cleaned = re.sub(r"(?s)<!--.*?-->", " ", cleaned)
    cleaned = re.sub(r"<[^>]+>", " ", cleaned)
    cleaned = _html.unescape(cleaned)
    cleaned = re.sub(r"[ \t]+", " ", cleaned)
    cleaned = re.sub(r"\n\s*\n+", "\n\n", cleaned).strip()
    return cleaned, meta


def parse_pdf(path: Path) -> tuple[str, dict]:
    """Extract text via pypdf/pypdf2/PyPDF2 when available, else byte-scan fallback."""
    meta: dict = {}
    text = ""
    for module_name in ("pypdf", "pypdf2", "PyPDF2"):
        try:
            module = __import__(module_name)
            reader_cls = getattr(module, "PdfReader", None)
            if reader_cls is None:
                continue
            reader = reader_cls(str(path))
            pages = []
            for page in getattr(reader, "pages", []):
                try:
                    pages.append(page.extract_text() or "")
                except Exception:
                    continue
            text = "\n\n".join(pages).strip()
            info = getattr(reader, "metadata", None)
            if info:
                for key in ("title", "author", "subject", "creator"):
                    value = getattr(info, key, None) or (info.get(f"/{key.capitalize()}", "") if hasattr(info, "get") else "")
                    if value:
                        meta[key] = str(value)
            if text:
                break
        except Exception:
            continue
    if not text:
        # Fallback: extract printable ASCII runs from the raw bytes.
        data = path.read_bytes()
        runs = re.findall(rb"[\x20-\x7e][\x20-\x7e\s]{20,}", data)
        decoded = []
        for run in runs:
            try:
                decoded.append(run.decode("ascii", errors="ignore"))
            except Exception:
                continue
        text = "\n".join(decoded)
        meta["parser_warning"] = "pdf library unavailable; used byte-scan fallback"
    if not meta.get("title"):
        first = next((line.strip() for line in text.splitlines() if line.strip()), "")
        if first:
            meta["title"] = first[:120]
    return text, meta


def parse_docx(path: Path) -> tuple[str, dict]:
    """Extract paragraphs from OOXML; falls back to zip/XML manual parse when
    python-docx is unavailable."""
    meta: dict = {}
    try:
        import docx  # type: ignore

        document = docx.Document(str(path))
        paragraphs = [p.text for p in document.paragraphs]
        text = "\n\n".join(paragraphs).strip()
        core = getattr(document, "core_properties", None)
        if core is not None:
            for key in ("title", "author", "subject", "keywords", "category"):
                value = getattr(core, key, "")
                if value:
                    meta[key] = str(value)
        if text:
            return text, meta
    except Exception:
        pass
    # Manual OOXML fallback (docx is a zip of XML files)
    try:
        with zipfile.ZipFile(path) as archive:
            try:
                core_xml = archive.read("docProps/core.xml")
                root = ET.fromstring(core_xml)
                for child in root.iter():
                    tag = child.tag.rsplit("}", 1)[-1]
                    if tag in {"title", "creator", "subject", "keywords"} and child.text:
                        meta[tag] = child.text.strip()
            except KeyError:
                pass
            doc_xml = archive.read("word/document.xml")
            root = ET.fromstring(doc_xml)
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            paragraphs = []
            for para in root.findall(".//w:p", ns):
                runs = [t.text or "" for t in para.findall(".//w:t", ns)]
                joined = "".join(runs).strip()
                if joined:
                    paragraphs.append(joined)
            text = "\n\n".join(paragraphs)
            return text, meta
    except Exception as exc:
        return "", {"parser_error": f"docx parse failed: {exc}"}
    return "", {"parser_error": "docx parse produced no text"}


PARSERS = {
    ".md": parse_markdown,
    ".markdown": parse_markdown,
    ".txt": parse_txt,
    ".html": parse_html,
    ".htm": parse_html,
    ".pdf": parse_pdf,
    ".docx": parse_docx,
}


def parse_file(path: Path, category: str = "general") -> NormalizedDocument:
    ext = path.suffix.lower()
    parser = PARSERS.get(ext)
    warnings: list[str] = []
    if parser is None:
        raise ValueError(f"Unsupported document type: {ext} ({path})")
    try:
        if ext in {".md", ".markdown"}:
            text, meta = parser(path)
        elif ext in {".html", ".htm"}:
            text, meta = parser(path)
        else:
            text, meta = parser(path)
    except Exception as exc:
        raise ValueError(f"Failed to parse {path}: {exc}") from exc
    if not text or not text.strip():
        warnings.append("empty document after parsing")
    title = str(meta.get("title") or path.stem.replace("_", " ").replace("-", " ").strip())
    doc_id = f"{category}:{path.stem}"
    if meta.get("parser_warning"):
        warnings.append(str(meta["parser_warning"]))
    if meta.get("parser_error"):
        warnings.append(str(meta["parser_error"]))
    return NormalizedDocument(
        doc_id=doc_id,
        title=title,
        text=text or "",
        source_path=str(path),
        extension=ext,
        category=category,
        raw_metadata=meta,
        warnings=warnings,
    )
